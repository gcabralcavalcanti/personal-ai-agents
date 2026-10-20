"""The habit export, the goals panel, the career page and the full demo run, on fictitious data."""

import io
import json
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from datetime import date
from pathlib import Path

from _path import DATA, ROOT
from openpyxl import load_workbook

from personal_ai_agents import pages
from personal_ai_agents.goals import panel
from personal_ai_agents.habits import dashboard, loop

END = date(2026, 11, 15)


class HabitExport(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.export = loop.write_fictitious_export(self.tmp / "export.zip", END, days=60)

    def test_export_has_the_app_layout(self):
        with zipfile.ZipFile(self.export) as z:
            names = z.namelist()
            head = z.read("Checkmarks.csv").decode().splitlines()
        self.assertIn("Habits.csv", names)
        self.assertTrue(head[0].startswith("Date,Plan tomorrow"))
        self.assertTrue(head[1].startswith(END.isoformat()))  # newest day first, as the app writes it

    def test_rates_ignore_days_not_recorded(self):
        marks = {END: "YES_MANUAL", date(2026, 11, 14): "NO", date(2026, 11, 13): "UNKNOWN", date(2026, 11, 12): "SKIP"}
        self.assertEqual(loop.rate(marks, END, 4), 0.5)
        self.assertIsNone(loop.rate({}, END, 7))
        self.assertEqual(
            loop.streak({END: "YES_MANUAL", date(2026, 11, 14): "SKIP", date(2026, 11, 13): "YES_AUTO"}, END), 2
        )

    def test_spreadsheet_and_page(self):
        habits, marks = loop.read_export(self.export)
        rows = dashboard.summary(habits, marks, END)
        dashboard.write_spreadsheet(self.tmp / "h.xlsx", habits, marks, rows, END, days=30)
        wb = load_workbook(self.tmp / "h.xlsx")
        self.assertEqual(wb.sheetnames, ["Summary", "Daily"])
        self.assertEqual(wb["Daily"].max_row, 31)
        html = dashboard.page(rows, marks, END, "export.zip", "h.xlsx")
        self.assertEqual(html.count('class="heat"'), len(habits))


class Goals(unittest.TestCase):
    def test_each_source_is_measured_and_scored(self):
        export = loop.write_fictitious_export(Path(tempfile.mkdtemp()) / "e.zip", END)
        _, marks = loop.read_export(export)
        goals = json.loads((Path(DATA) / "goals.json").read_text(encoding="utf-8"))
        results = panel.evaluate(goals, marks, END)
        kinds = {k["source"] for o in results for k in o["krs"]}
        self.assertEqual(kinds, {"artifact", "habit app", "task manager"})
        for o in results:
            expected = round(10 * sum(panel.CREDIT[k["status"]] for k in o["krs"]) / len(o["krs"]))
            self.assertEqual(o["measured"], expected)
        self.assertIn("felt", panel.page(goals, results, END))


class FullDemo(unittest.TestCase):
    def test_all_pages_written_and_index_links_them(self):
        import demo

        out = io.StringIO()
        with redirect_stdout(out):
            self.assertEqual(demo.main(["--no-open"]), 0)
        folder = Path(ROOT) / "output"
        for name, *_ in pages.AGENTS:
            self.assertTrue((folder / f"{name}.html").exists(), name)
        index = (folder / "index.html").read_text(encoding="utf-8")
        self.assertTrue(all(f'href="{name}.html"' in index for name, *_ in pages.AGENTS))
        career = (folder / "career.html").read_text(encoding="utf-8")
        self.assertIn("decided by code", career)


if __name__ == "__main__":
    unittest.main()
